"""Excel (xlsx/xlsm — không bao giờ chạy macro), CSV/TSV, DOCX, JSON, TXT, ảnh."""

import csv
import io
import json
from pathlib import Path

from .. import settings
from ..ingest.sniff import decode_text
from . import docai, ocr, tables, text_invoice
from .base import Doc, ExtractError, Result

VERSION = "1.3.0"  # 1.1: bảng bất kỳ; 1.2: tổng tiền bằng chữ; 1.3: PaddleOCR dự phòng cho ảnh


def _trim(grid):
    width = 0
    for row in grid:
        for i in range(len(row) - 1, -1, -1):
            if row[i] not in (None, ""):
                width = max(width, i + 1)
                break
    return [list(row[:width]) for row in grid]


def excel(path, file_name):
    import openpyxl

    max_rows = int(settings.get()["max_sheet_rows"])
    res = Result("excel_openpyxl", VERSION)
    try:
        # Bản gốc lưu theo SHA-256 (không có đuôi file) → truyền nội dung, openpyxl không kiểm tra đuôi
        data = Path(path).read_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False)
        wbf = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_vba=False)
    except Exception as e:
        raise ExtractError("bad_xlsx", f"Không đọc được workbook: {type(e).__name__}") from e
    texts, sheets = [], []
    for ws in wb.worksheets:
        grid = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                res.status = "PARTIAL"
                break
            grid.append(list(row))
        grid = _trim(grid)
        uncached = 0
        try:
            for frow, vrow in zip(wbf[ws.title].iter_rows(values_only=True), grid):
                for fv, vv in zip(frow, vrow):
                    if isinstance(fv, str) and fv.startswith("=") and vv is None:
                        uncached += 1
        except Exception:
            pass
        docs = tables.build_docs(grid, ws.title, file_name)
        for d in docs:
            if uncached:
                d.warnings.append(f"{uncached} ô công thức chưa có giá trị đã tính (file chưa từng được Excel tính lại)")
            res.docs.append(d)
        sheets.append({"name": ws.title, "rows": len(grid), "tables": len(docs), "uncached_formulas": uncached})
        texts.append("\n".join(" | ".join(str(v) for v in r if v not in (None, "")) for r in grid[:2000]))
    wb.close()
    wbf.close()
    res.text = "\n".join(texts)
    res.stats = {"sheets": sheets}
    if not res.docs:
        res.docs.append(Doc(doc_type="text", title=file_name, data={"sheets": sheets},
                            warnings=["Không nhận diện được bảng dữ liệu có tiêu đề"]))
    return res


def xls(data, file_name):
    """Excel 97-2003 (.xls) qua xlrd: chỉ đọc giá trị đã lưu; không phân tích được công thức/macro."""
    import xlrd

    max_rows = int(settings.get()["max_sheet_rows"])
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as e:
        raise ExtractError("bad_xls", f"Không đọc được file .xls: {type(e).__name__}") from e
    res = Result("excel_xlrd", VERSION)
    texts, sheets = [], []
    for sh in book.sheets():
        grid = []
        for r in range(min(sh.nrows, max_rows)):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:
                        row.append(None)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    row.append(int(cell.value))
                else:
                    row.append(cell.value)
            grid.append(row)
        if sh.nrows > max_rows:
            res.status = "PARTIAL"
        grid = _trim(grid)
        docs = tables.build_docs(grid, sh.name, file_name)
        for d in docs:
            d.warnings.append("File .xls: chỉ đọc giá trị đã lưu, không kiểm tra được công thức")
        res.docs.extend(docs)
        sheets.append({"name": sh.name, "rows": len(grid), "tables": len(docs)})
        texts.append("\n".join(" | ".join(str(v) for v in r if v not in (None, "")) for r in grid[:2000]))
    book.release_resources()
    res.text = "\n".join(texts)
    res.stats = {"sheets": sheets}
    if not res.docs:
        res.docs.append(Doc(doc_type="text", title=file_name, data={"sheets": sheets},
                            warnings=["Không nhận diện được bảng dữ liệu có tiêu đề"]))
    return res


def delimited(data, file_name, kind):
    text, enc = decode_text(data)
    if text is None:
        raise ExtractError("bad_encoding", "Không giải mã được văn bản (encoding không hỗ trợ)")
    delim = "\t" if kind == "tsv" else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff("\n".join(text.splitlines()[:50]), delimiters=",;|\t").delimiter
        except csv.Error:
            delim = ","
    grid = _trim([row for row in csv.reader(io.StringIO(text), delimiter=delim)])
    max_rows = int(settings.get()["max_sheet_rows"])
    res = Result("delimited", VERSION, stats={"encoding": enc, "delimiter": delim, "rows": len(grid)})
    if len(grid) > max_rows:
        grid, res.status = grid[:max_rows], "PARTIAL"
    sheet = "CSV" if kind == "csv" else "TSV"
    res.docs = tables.build_docs(grid, sheet, file_name) or [Doc(doc_type="text", title=file_name)]
    res.text = text[:200000]
    return res


def docx(data, file_name):
    import docx as pydocx

    try:
        d = pydocx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError("bad_docx", f"Không đọc được DOCX: {type(e).__name__}") from e
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    res = Result("docx", VERSION)
    for ti, t in enumerate(d.tables, start=1):
        grid = [[c.text.strip() for c in row.cells] for row in t.rows]
        res.docs.extend(tables.build_docs(grid, f"Bảng {ti}", file_name))
        paras.extend(" | ".join(r) for r in grid)
    res.text = "\n".join(paras)
    if text_invoice.looks_like_invoice(res.text):
        inv = text_invoice.extract([(1, [(p, None, None) for p in paras])], file_name, "docx_text")
        if inv:
            res.docs.insert(0, inv)
    if not res.docs:
        res.docs.append(Doc(doc_type="text", title=file_name))
    res.stats = {"paragraphs": len(paras), "tables": len(d.tables)}
    return res


def jsonfile(data, file_name):
    text, _ = decode_text(data)
    try:
        obj = json.loads(text or "")
    except ValueError as e:
        raise ExtractError("bad_json", "JSON không hợp lệ") from e
    res = Result("json", VERSION, text=(text or "")[:200000])
    rows = obj if isinstance(obj, list) else next((v for v in obj.values() if isinstance(v, list)), None) if isinstance(obj, dict) else None
    if rows and all(isinstance(r, dict) for r in rows):
        keys = []
        for r in rows:
            for k in r:
                if k not in keys:
                    keys.append(k)
        grid = [keys] + [[r.get(k) for k in keys] for r in rows]
        res.docs = tables.build_docs(grid, "JSON", file_name)
    if not res.docs:
        res.docs.append(Doc(doc_type="text", title=file_name))
    return res


def textfile(data, file_name):
    text, enc = decode_text(data)
    if text is None:
        raise ExtractError("bad_encoding", "Không giải mã được văn bản")
    res = Result("text", VERSION, text=text[:500000], stats={"encoding": enc})
    if text_invoice.looks_like_invoice(text):
        inv = text_invoice.extract([(1, [(ln, None, None) for ln in text.splitlines()])], file_name, "plain_text")
        if inv:
            res.docs.append(inv)
    if not res.docs:
        res.docs.append(Doc(doc_type="text", title=file_name))
    return res


def image(data, file_name):
    paddle_ok = docai.available("ocr") and settings.get()["ocr"] != "off"
    if not ocr.available() and not paddle_ok:
        res = Result("image_ocr", VERSION, status="PARTIAL")
        res.docs.append(Doc(doc_type="text", title=file_name, status="NEEDS_REVIEW",
                            warnings=["Ảnh chứng từ cần OCR nhưng OCR đang tắt/không khả dụng"]))
        return res
    import numpy as np
    from PIL import Image

    try:
        im = Image.open(io.BytesIO(data))
        im.seek(0)  # TIFF nhiều trang: trang đầu
        im = im.convert("RGB")
    except Exception as e:
        raise ExtractError("bad_image", f"Ảnh hỏng: {type(e).__name__}") from e
    if im.width * im.height > 40_000_000:
        raise ExtractError("image_too_large", "Ảnh quá lớn (> 40 megapixel)")
    lines, model = (ocr.read_image(np.asarray(im)), ocr.MODEL_ID) if ocr.available() else ([], None)
    text = "\n".join(t for t, _, _ in lines)
    inv = text_invoice.extract([(1, lines)], file_name, "ocr_image", base_conf=0.75) if text_invoice.looks_like_invoice(text) else None
    if inv is None and paddle_ok:  # dự phòng PaddleOCR (ảnh mờ / không có EasyOCR)
        buf = io.BytesIO()
        im.save(buf, "PNG")
        alt, alt_model = docai.ocr_image(buf.getvalue())
        if alt is not None:
            alt_text = "\n".join(t for t, _, _ in alt)
            alt_inv = text_invoice.extract([(1, alt)], file_name, "ocr_image", base_conf=0.7) if text_invoice.looks_like_invoice(alt_text) else None
            if alt_inv is not None or not lines:
                lines, text, inv, model = alt, alt_text, alt_inv, alt_model
    res = Result("image_ocr", VERSION, model_id=model, text=text,
                 stats={"width": im.width, "height": im.height, "lines": len(lines)})
    res.docs.append(inv or Doc(doc_type="text", title=file_name, status="NEEDS_REVIEW",
                               warnings=["Văn bản từ OCR — cần kiểm duyệt"]))
    return res
