"""
Phân tích workbook Excel (không chạy macro, không tính lại toàn bộ công thức).

Phân biệt ba loại giá trị: công thức (text), giá trị đã lưu trong file (cached),
và giá trị tính lại. Hệ thống chỉ TỰ TÍNH LẠI được công thức SUM(phạm vi) và cộng/trừ ô đơn giản;
các công thức khác chỉ được kiểm tra cấu trúc. Kết quả nói rõ phần nào đã/không tính lại.
"""

import io
import re
from pathlib import Path

from .. import money
from ..extract.base import fold

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
SUM_RE = re.compile(r"^=\s*SUM\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$", re.I)
ADD_RE = re.compile(r"^=\s*([+-]?\s*\$?[A-Z]{1,3}\$?\d+(\s*[+-]\s*\$?[A-Z]{1,3}\$?\d+)*)\s*$", re.I)
CELL_RE = re.compile(r"([+-]?)\s*\$?([A-Z]{1,3})\$?(\d+)", re.I)
TOTAL_WORDS = ("tong", "cong", "total", "sum", "tong cong")


def col_index(letters):
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return money.to_decimal(v)
    return None


def analyze(path, file_name, max_cells=500_000):
    import openpyxl

    data = Path(path).read_bytes()
    wbf = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=False, keep_links=True)
    wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_vba=False, keep_links=True)
    out = {"file": file_name, "sheets": [], "findings": [], "defined_names": [], "external_links": 0,
           "recalc_scope": "Chỉ tính lại SUM(phạm vi) và cộng/trừ ô đơn giản; công thức khác chỉ kiểm tra cấu trúc"}
    try:
        out["defined_names"] = [{"name": n, "value": str(dn.attr_text)} for n, dn in wbf.defined_names.items()]
    except Exception:
        pass
    out["external_links"] = len(getattr(wbf, "_external_links", []) or [])
    if out["external_links"]:
        out["findings"].append({"severity": "warning", "sheet": None, "cell": None, "type": "external_links",
                                "message": f"Workbook có {out['external_links']} liên kết ngoài — giá trị phụ thuộc file khác"})
    cells = 0
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        info = {"name": ws.title, "dimensions": ws.dimensions, "formulas": 0, "recalculated": 0,
                "uncached": 0, "merged": [str(r) for r in ws.merged_cells.ranges][:200],
                "tables": list(ws.tables.keys()) if hasattr(ws, "tables") else [],
                "charts": len(getattr(ws, "_charts", [])), "pivots": len(getattr(ws, "_pivots", [])),
                "data_validations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
                "state": ws.sheet_state}
        if ws.sheet_state != "visible":
            out["findings"].append({"severity": "info", "sheet": ws.title, "cell": None, "type": "hidden_sheet",
                                    "message": f"Sheet '{ws.title}' đang ẩn ({ws.sheet_state})"})
        for row in ws.iter_rows():
            for c in row:
                cells += 1
                if cells > max_cells:
                    out["findings"].append({"severity": "warning", "sheet": ws.title, "cell": None, "type": "truncated",
                                            "message": "Workbook quá lớn — chỉ phân tích một phần"})
                    break
                v = c.value
                cached = wv[c.coordinate].value
                if isinstance(cached, str) and cached in ERRORS:
                    out["findings"].append({"severity": "error", "sheet": ws.title, "cell": c.coordinate,
                                            "type": "error_value", "message": f"Ô {c.coordinate} có lỗi {cached}",
                                            "formula": v if isinstance(v, str) else None})
                if not (isinstance(v, str) and v.startswith("=")):
                    if _num(v) is not None and _label_is_total(ws, c):
                        out["findings"].append({"severity": "warning", "sheet": ws.title, "cell": c.coordinate,
                                                "type": "hardcoded_total",
                                                "message": f"Ô tổng {c.coordinate} là số nhập tay, không phải công thức"})
                    continue
                info["formulas"] += 1
                if "#REF!" in v.upper():
                    out["findings"].append({"severity": "error", "sheet": ws.title, "cell": c.coordinate,
                                            "type": "broken_reference", "message": f"Công thức {c.coordinate} tham chiếu ô đã bị xoá (#REF!)",
                                            "formula": v})
                if re.search(r"\[\d+\]|\.xlsx?\]", v, re.I):
                    out["findings"].append({"severity": "warning", "sheet": ws.title, "cell": c.coordinate,
                                            "type": "external_reference", "message": f"Công thức {c.coordinate} tham chiếu workbook khác",
                                            "formula": v})
                if cached is None:
                    info["uncached"] += 1
                recalc = _recalc(v, wv)
                if recalc is None:
                    continue
                info["recalculated"] += 1
                value, ref = recalc
                cv = _num(cached)
                if cv is not None and abs(cv - value) > money.to_decimal("0.000001"):
                    out["findings"].append({"severity": "error", "sheet": ws.title, "cell": c.coordinate,
                                            "type": "stale_or_wrong_value",
                                            "message": f"Giá trị lưu {money.fmt(cv)} ≠ tính lại {money.fmt(value)}",
                                            "formula": v})
                m = SUM_RE.match(v)
                if m:
                    _check_sum_range(ws, wv, c, m, out)
        out["sheets"].append(info)
    wbf.close()
    wbv.close()
    out["summary"] = {
        "errors": sum(1 for f in out["findings"] if f["severity"] == "error"),
        "warnings": sum(1 for f in out["findings"] if f["severity"] == "warning"),
        "formulas": sum(s["formulas"] for s in out["sheets"]),
        "recalculated": sum(s["recalculated"] for s in out["sheets"]),
        "uncached": sum(s["uncached"] for s in out["sheets"]),
    }
    return out


def _label_is_total(ws, c):
    """Ô số nằm cùng dòng với nhãn 'Tổng/Cộng/Total' ở các cột bên trái."""
    for col in range(1, c.column):
        lab = ws.cell(row=c.row, column=col).value
        if isinstance(lab, str) and fold(lab).startswith(TOTAL_WORDS):
            return True
    return False


def _recalc(formula, wv):
    m = SUM_RE.match(formula)
    if m:
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        if (r2 - r1 + 1) * (col_index(c2) - col_index(c1) + 1) > 100_000:
            return None
        total = money.ZERO
        for r in range(r1, r2 + 1):
            for ci in range(col_index(c1), col_index(c2) + 1):
                val = wv.cell(row=r, column=ci).value
                if isinstance(val, str) and val.startswith("="):
                    return None
                n = _num(val)
                if n is not None:
                    total += n
        return total, f"{c1}{r1}:{c2}{r2}"
    m = ADD_RE.match(formula)
    if m:
        total = money.ZERO
        for sign, col, row in CELL_RE.findall(m.group(1)):
            val = wv[f"{col.upper()}{row}"].value
            n = _num(val)
            if n is None:
                return None
            total += -n if sign == "-" else n
        return total, m.group(1)
    return None


def _check_sum_range(ws, wv, c, m, out):
    """SUM theo cột kết thúc ngay trên ô tổng nhưng bỏ sót dòng số liền kề phía trên/dưới phạm vi."""
    c1, r1, c2, r2 = m.group(1).upper(), int(m.group(2)), m.group(3).upper(), int(m.group(4))
    if c1 != c2 or c.column != col_index(c1):
        return
    above = r1 - 1
    if above >= 1 and _num(wv.cell(row=above, column=c.column).value) is not None:
        out["findings"].append({"severity": "warning", "sheet": ws.title, "cell": c.coordinate, "type": "sum_range_gap",
                                "message": f"SUM({c1}{r1}:{c2}{r2}) bỏ qua ô số {c1}{above} ngay phía trên phạm vi"})
    below = r2 + 1
    if below < c.row and _num(wv.cell(row=below, column=c.column).value) is not None:
        out["findings"].append({"severity": "error", "sheet": ws.title, "cell": c.coordinate, "type": "sum_range_gap",
                                "message": f"SUM({c1}{r1}:{c2}{r2}) bỏ sót {c1}{below}…{c1}{c.row - 1} trước ô tổng"})
