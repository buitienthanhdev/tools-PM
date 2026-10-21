"""
AI Tư vấn pháp luật — RAG.

Nạp văn bản → làm sạch + tách theo Chương/Mục/Điều/Khoản → chỉ mục lai (BM25 + vector Ollama, hợp nhất RRF)
→ (câu hỏi nối tiếp được viết lại thành câu độc lập) → AI local trả lời có đánh số căn cứ [n]
→ kiểm chứng: có trích dẫn không, có nhắc Điều nào ngoài các căn cứ đã truy xuất không.
Dưới ngưỡng tương đồng thì nói thẳng "chưa có căn cứ" thay vì để AI đoán.
"""

import hashlib
import json
import os
import re
import threading
import time
import uuid

from core import config as hubcfg
from core import llm

from . import index as ix
from . import parse as ps
from . import vecstore

KB_DIR = os.path.join(hubcfg.DATA_DIR, "legal")  # PM_DATA_DIR đổi được khi thử nghiệm
KB_PATH = os.path.join(KB_DIR, "kb.json")

_lock = threading.RLock()
_kb = {"docs": []}
_index = None
parse_law = ps.parse_law  # giữ tên cũ cho nơi khác import


# ── Lưu trữ ──────────────────────────────────────────────────
def _doc_hash(d):
    return hashlib.sha1("\n".join(c["text"] for c in d["chunks"]).encode()).hexdigest()


def load():
    global _kb, _index
    with _lock:
        if os.path.exists(KB_PATH):
            with open(KB_PATH, encoding="utf-8") as f:
                _kb = json.load(f)
        seen, kept = set(), []
        for d in _kb["docs"]:  # bỏ văn bản trùng nội dung
            h = _doc_hash(d)
            if h not in seen:
                seen.add(h)
                kept.append(d)
        if len(kept) != len(_kb["docs"]):
            print(f"[legal_rag] bỏ {len(_kb['docs']) - len(kept)} văn bản trùng nội dung")
            _kb["docs"] = kept
            _save()
        _index = None


def _save():
    os.makedirs(KB_DIR, exist_ok=True)
    tmp = KB_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(_kb, f, ensure_ascii=False)
    os.replace(tmp, KB_PATH)


def _reset():
    global _index
    _index = None


# ── Đọc file ─────────────────────────────────────────────────
_ocr_reader = None


def _ocr_pdf(data: bytes, progress=None) -> str:
    """OCR cho PDF scan (easyocr vi+en, GPU nếu có)."""
    global _ocr_reader
    import numpy as np
    import pymupdf

    if _ocr_reader is None:
        import easyocr

        _ocr_reader = easyocr.Reader(["vi", "en"], gpu=True, verbose=False)
    doc = pymupdf.open(stream=data, filetype="pdf")
    pages = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=200, colorspace=pymupdf.csGRAY)
        img = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w)
        pages.append("\n".join(_ocr_reader.readtext(img, detail=0, paragraph=False)))
        if progress:
            progress(i + 1, len(doc))
    return "\n".join(pages)


def extract_text(filename: str, data: bytes, progress=None) -> str:
    import io

    ext = os.path.splitext(filename)[1].lower()
    if ext in (".txt", ".md", ""):
        for enc in ("utf-8-sig", "utf-16", "cp1258"):
            try:
                return data.decode(enc)
            except UnicodeError:
                continue
        return data.decode("utf-8", errors="ignore")
    if ext == ".docx":
        import docx  # python-docx

        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:  # bảng (vd. phụ lục) cũng là nội dung luật
            lines += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
        return "\n".join(lines)
    if ext == ".pdf":
        try:  # pymupdf giữ thứ tự đọc và dấu tiếng Việt tốt hơn pypdf
            import pymupdf

            doc = pymupdf.open(stream=data, filetype="pdf")
            text = "\n".join(p.get_text("text") for p in doc)
        except ImportError:
            from pypdf import PdfReader

            text = "\n".join(
                (p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages
            )
        if len(text.strip()) < 200:  # PDF scan → không có lớp chữ
            text = _ocr_pdf(data, progress)
        return text
    raise ValueError(f"Định dạng {ext} chưa được hỗ trợ (dùng .txt .md .docx .pdf)")


# ── Quản lý văn bản ──────────────────────────────────────────
def add_law(name: str, text: str, save: bool = True):
    text = ps.clean_text(text)
    chunks = ps.parse_law(text)
    if not chunks:
        raise ValueError("Không tìm thấy nội dung để nạp")
    doc = {
        "id": uuid.uuid4().hex[:10],
        "name": name.strip(),
        "added": int(time.time()),
        "chunks": chunks,
        "quality": ps.quality_report(text, chunks),
    }
    with _lock:
        # cùng tên = cập nhật bản mới
        _kb["docs"] = [
            d for d in _kb["docs"] if d["name"].lower() != doc["name"].lower()
        ]
        h = _doc_hash(doc)
        if any(_doc_hash(d) == h for d in _kb["docs"]):
            raise ValueError("Nội dung này đã có trong hệ thống (trùng với văn bản khác)")
        _kb["docs"].append(doc)
        if save:  # nạp hàng loạt: save=False rồi gọi flush() một lần ở cuối
            _save()
            _reset()
    return doc


def flush():
    with _lock:
        _save()
        _reset()


def delete_law(doc_id: str) -> bool:
    with _lock:
        before = len(_kb["docs"])
        _kb["docs"] = [d for d in _kb["docs"] if d["id"] != doc_id]
        if len(_kb["docs"]) != before:
            _save()
            _reset()
            return True
    return False


def list_laws():
    return [
        {
            "id": d["id"],
            "name": d["name"],
            "added": d["added"],
            "articles": len(d["chunks"]),
            "warnings": (d.get("quality") or {}).get("warnings", []),
            "superseded_by": d.get("superseded_by"),
        }
        for d in _kb["docs"]
    ]


# ── Chỉ mục ──────────────────────────────────────────────────
def _embed_fn():
    model = hubcfg.model("embed")
    return lambda texts: llm.embed(model, texts, name="legal-embed")


# ── Văn bản bị thay thế / bãi bỏ ─────────────────────────────
_CODE = re.compile(r"(\d{1,4})\s*/\s*(\d{4})(?:\s*/\s*([\wĐđ\-]+))?")
_REPLACE = re.compile(
    r"(?:thay thế|bãi bỏ|thay thế cho)\s+(?:toàn bộ\s+)?(Nghị định|Thông tư(?: liên tịch)?|Luật|Bộ luật|Quyết định)\s+(?:số\s+)?"
    r"(\d{1,4})\s*/\s*(\d{4})(?:\s*/\s*([\wĐđ\-]+))?",
    re.I,
)


def _kind(word: str) -> str:
    w = word.lower()
    return "nd" if w.startswith("nghị") else "tt" if w.startswith("thông") else "luat" if w in ("luật", "bộ luật") else "qd"


def _issuer(code: str) -> str:
    """'TT-BTC' → 'BTC'; 'NĐ-CP' → 'CP'; 'QH14' → ''."""
    c = (code or "").upper().replace("Đ", "D")
    return c.split("-")[-1] if c.startswith(("TT", "ND", "QD")) and "-" in c else ""


def _doc_key(name: str):
    """'Nghị định 98/2020/NĐ-CP ...' → ('nd', '98/2020', 'CP'); 'Luật X (59/2020/QH14)' → ('luat', '59/2020', '')."""
    m = _CODE.search(name)
    if not m:
        return None
    first = name.strip().lower()
    kind = "luat" if first.startswith(("luật", "bộ luật")) else _kind(first.split(" ")[0])
    return (kind, f"{int(m.group(1))}/{m.group(2)}", _issuer(m.group(3)) if kind != "luat" else "")


def mark_superseded():
    """Đọc điều khoản thi hành: 'thay thế / bãi bỏ <loại> số N/YYYY[/cơ quan]' → đánh dấu văn bản đó hết hiệu lực
    (chỉ khi cùng loại, đúng cơ quan ban hành nếu có nêu, văn bản thay thế không cũ hơn). Văn bản bị đánh dấu không đưa vào chỉ mục."""
    by_key = {}
    for d in _kb["docs"]:
        k = _doc_key(d["name"])
        if k:
            by_key.setdefault(k[:2], []).append((k, d))
    for d in _kb["docs"]:
        d.pop("superseded_by", None)
    found = 0
    for d in _kb["docs"]:
        me = _doc_key(d["name"])
        if not me:
            continue
        text = " ".join(c["text"] for c in d["chunks"][-40:])  # điều khoản thi hành nằm cuối văn bản
        for m in _REPLACE.finditer(text):
            kind, key = _kind(m.group(1)), f"{int(m.group(2))}/{m.group(3)}"
            issuer = _issuer(m.group(4)) if kind != "luat" else ""
            cands = [(k, old) for k, old in by_key.get((kind, key), []) if old is not d]
            if issuer:
                cands = [(k, old) for k, old in cands if k[2] == issuer]
            elif len(cands) != 1:
                continue  # tham chiếu thiếu cơ quan ban hành mà có nhiều ứng viên: không đoán
            if int(m.group(3)) > int(me[1].split("/")[1]):
                continue
            for _, old in cands:
                if not old.get("superseded_by"):
                    old["superseded_by"] = d["name"]
                    found += 1
    with _lock:
        _save()
        _reset()
    return found


def get_index(rebuild=False):
    """Chỉ mục xây khi cần (vector tính một lần rồi lưu cache theo nội dung)."""
    global _index
    with _lock:
        if _index is None or rebuild:
            _index = ix.Index(
                [d for d in _kb["docs"] if not d.get("superseded_by")], KB_DIR, embed_fn=_embed_fn(), embed_model=hubcfg.model("embed")
            )
        return _index


def reindex():
    """Xoá cache vector và tính lại (dùng khi đổi model embedding)."""
    vecstore.clear(KB_DIR)
    idx = get_index(rebuild=True)
    return {"chunks": len(idx.entries), "vectors": idx.vecs is not None, "error": idx.embed_error}


def index_status():
    idx = _index
    return {
        "chunks": sum(len(d["chunks"]) for d in _kb["docs"]),
        "vectors": bool(idx and idx.vecs is not None),
        "embed_model": hubcfg.model("embed"),
        "embed_error": idx.embed_error if idx else None,
    }


# ── Truy xuất ────────────────────────────────────────────────
_FOLLOWUP = re.compile(r"\b(nó|vậy|thế|đó|này|còn|trường hợp đó|như trên|thì sao)\b", re.I)


def rewrite_query(question, history):
    """Câu hỏi nối tiếp ("vậy còn nghỉ thai sản?") → câu hỏi độc lập để tìm kiếm."""
    turns = [
        m for m in (history or [])[-4:]
        if m.get("role") in ("user", "assistant") and m.get("content")
    ]
    if not turns or (len(question.split()) > 14 and not _FOLLOWUP.search(question)):
        return question
    convo = "\n".join(f'{m["role"]}: {str(m["content"])[:400]}' for m in turns)
    try:
        out = llm.chat(
            hubcfg.model("legal"),
            [
                {
                    "role": "system",
                    "content": "Viết lại câu hỏi cuối của người dùng thành MỘT câu hỏi đầy đủ, độc lập, "
                    "giữ đủ chủ thể/tình huống nhắc ở các lượt trước. Chỉ trả về câu hỏi, không giải thích.",
                },
                {"role": "user", "content": f"{convo}\nuser: {question}"},
            ],
            num_ctx=4096, num_predict=120, temperature=0.0, timeout=120, name="legal-rewrite",
        )
        out = out.strip().strip('"')
        return out if 5 <= len(out) <= 400 else question
    except Exception as e:
        print(f"[legal_rag] rewrite lỗi: {e}")
        return question


def search(question: str, k=6, history=None):
    """Trả về danh sách đoạn luật đã hợp nhất, có điểm BM25 + cosine."""
    return retrieve(question, k=k, history=history)["hits"]


def retrieve(question: str, k=6, history=None, mode="auto"):
    """mode: auto (lai nếu có vector) | bm25 (ép chỉ BM25, dùng để so sánh chất lượng)."""
    idx = get_index()
    if not idx.entries:
        return {"hits": [], "query": question, "mode": "empty", "confidence": 0.0}
    query = rewrite_query(question, history)
    qvec = None
    if idx.vecs is not None and mode != "bm25":
        try:
            qvec = llm.embed(
                hubcfg.model("embed"), [ix.QUERY_INSTRUCT + query], name="legal-q"
            )[0]
        except Exception as e:
            print(f"[legal_rag] embed câu hỏi lỗi, dùng BM25: {e}")
    accented = ix.has_diacritics(query)
    # không dấu → vector kém tin cậy: giảm trọng số, BM25 (có bản bỏ dấu) gánh chính
    results = idx.search(query, qvec, k=k, dense_weight=1.0 if accented else 0.1)
    hits = [
        {
            "law": e["doc"], "article": e["chunk"].get("article", ""),
            "title": e["chunk"].get("title", ""), "clause": e["chunk"].get("clause", ""),
            "chapter": e["chunk"].get("chapter", ""), "section": e["chunk"].get("section", ""),
            "text": e["chunk"]["text"], **info,
        }
        for e, info in results
    ]
    dense = [h["dense"] for h in hits if h["dense"] is not None]
    return {
        "hits": hits,
        "query": query,
        "mode": "hybrid" if qvec is not None else "bm25",
        "confidence": round(max(dense), 3) if dense and accented else None,
    }


# ── Trả lời ──────────────────────────────────────────────────
SYSTEM_PROMPT = """Bạn là trợ lý pháp luật Việt Nam của PM Tools: thân thiện, tích cực, nói chuyện tự nhiên như một người bạn am hiểu luật.

CÁCH TRẢ LỜI
- Câu chào, cảm ơn, hỏi thăm, trò chuyện thường: đáp ngắn gọn, vui vẻ, rồi nhẹ nhàng mời người dùng nêu vấn đề cần hỏi. Không tra cứu, không trích dẫn.
- Câu hỏi pháp luật: trả lời thẳng vào ý chính ngay câu đầu, rồi giải thích dễ hiểu bằng lời thường (điều kiện, ngoại lệ, mức phạt/quyền/nghĩa vụ nếu có). Đi theo mạch câu hỏi, độ dài vừa đủ: hỏi ngắn thì đáp gọn, tình huống phức tạp mới trình bày kỹ hơn. Chỉ dùng gạch đầu dòng khi thực sự có nhiều ý song song.
- Mỗi lần diễn đạt khác nhau: không dùng một câu mở đầu cố định, không lặp công thức "Cảm ơn câu hỏi của bạn", không rập khuôn.
- Giọng ấm áp, tích cực, trấn an khi người dùng đang lo lắng; nếu có hướng giải quyết thì nêu ra.
- Không nhắc lại đoạn nhắc "tham khảo, hãy hỏi luật sư" ở mỗi câu. Chỉ khuyên gặp luật sư khi vụ việc phức tạp, tranh chấp hoặc rủi ro lớn, và nói một câu ngắn, tự nhiên.

NGUYÊN TẮC VỀ CĂN CỨ (bắt buộc)
- Khi kết luận về luật, CHỈ dựa vào phần "CĂN CỨ TRUY XUẤT". Không bịa điều, khoản, mức phạt, số liệu; không nhắc Điều nào không có trong căn cứ.
- Mỗi nhận định lấy từ căn cứ kết thúc bằng số căn cứ trong ngoặc vuông, ví dụ "... không quá 12 tháng [2]." hoặc [1][3]. Có thể nói tự nhiên "theo Điều 8 của Luật ..." rồi kèm [n].
- Nếu căn cứ không liên quan hoặc không đủ: nói thật và nhẹ nhàng rằng hiện chưa tìm thấy quy định phù hợp trong dữ liệu, KHÔNG đoán và KHÔNG tự nêu quy định pháp luật từ trí nhớ. Hãy gợi ý cách hỏi cụ thể hơn (nêu hành vi, lĩnh vực, đối tượng) hoặc hướng tìm nguồn chính thức. Trường hợp này không gắn [n].
- Không nhắc tới "căn cứ truy xuất", "cơ sở dữ liệu" hay cách hệ thống hoạt động trừ khi cần giải thích vì sao chưa trả lời được.
- Người dùng không hề cung cấp căn cứ hay văn bản nào, đó là dữ liệu của hệ thống: không nói "bạn cung cấp/bạn gửi" và không yêu cầu họ gửi văn bản luật. Khi chưa có quy định phù hợp, nói kiểu "mình chưa tìm thấy quy định về việc này trong kho luật hiện có".
- Trả lời bằng tiếng Việt."""

NO_BASIS_HINT = (
    "CĂN CỨ TRUY XUẤT: (không tìm thấy đoạn luật nào đủ liên quan trong dữ liệu hiện có)\n"
    "Nếu đây là câu chào/trò chuyện thì cứ đáp tự nhiên. Nếu là câu hỏi pháp luật thì nói thật là chưa có "
    "quy định phù hợp trong dữ liệu, gợi ý cách hỏi cụ thể hơn, tuyệt đối không tự nêu nội dung luật."
)


def _label(h):
    ref = f'Điều {h["article"]}. {h["title"]}'.strip(" .") if h["article"] else "Trích đoạn"
    if h.get("clause"):
        ref += f' (khoản {h["clause"]})'
    return f'{h["law"]} — {ref}'


def _context(hits):
    return "\n\n".join(
        f"[{i}] {_label(h)}"
        + (f' · {h["chapter"]}' if h.get("chapter") else "")
        + f'\n{h["text"][:2500]}'
        for i, h in enumerate(hits, 1)
    )


def _call_llm(question, hits, history):
    """Trả lời bằng AI local (Ollama) dựa trên các căn cứ đã truy xuất."""
    msgs = [{"role": "system", "content": SYSTEM_PROMPT}]
    for m in (history or [])[-6:]:
        if m.get("role") in ("user", "assistant") and m.get("content"):
            msgs.append({"role": m["role"], "content": str(m["content"])[:3000]})
    ctx = f"CĂN CỨ TRUY XUẤT:\n{_context(hits)}" if hits else NO_BASIS_HINT
    msgs.append({"role": "user", "content": f"{ctx}\n\nNGƯỜI DÙNG: {question}"})
    return llm.chat(
        hubcfg.model("legal"), msgs, num_ctx=12288, num_predict=1500,
        temperature=0.6, name="legal", timeout=600,
    )


def _fallback(hits):
    top = hits[0]
    return (
        f"Mình chưa soạn được câu trả lời chi tiết lúc này, nhưng quy định gần nhất với câu hỏi của bạn là "
        f'{_label(top)}:\n\n{top["text"][:1500]}'
    )


def verify_answer(text, hits):
    """Kiểm chứng câu trả lời với căn cứ đã truy xuất."""
    cited = sorted(
        {int(n) for n in re.findall(r"\[(\d{1,2})\]", text) if 1 <= int(n) <= len(hits)}
    )
    known = {h["article"] for h in hits if h["article"]}
    mentioned = set(re.findall(r"Điều\s+(\d+[a-zA-Z]?)", text))
    unverified = sorted(mentioned - known)
    return {
        "cited": cited,
        "unverified_articles": unverified,
        "grounded": bool(cited) and not unverified,
    }


def answer(question: str, history=None):
    history = history or []
    try:
        r = retrieve(question, k=6, history=history)
    except llm.LLMError as e:
        return {"answer": f"Không tìm kiếm được căn cứ: {e}", "sources": [],
                "grounded": False, "ai_error": str(e)}
    hits = r["hits"]
    min_sim = float(hubcfg.get().get("legal_min_similarity", 0.35))
    weak = r["confidence"] is not None and r["confidence"] < min_sim
    if weak:
        hits = []  # không đủ căn cứ: để AI đáp tự nhiên, không đính căn cứ không liên quan
    ai_error, text = None, None
    try:
        text = _call_llm(question, hits, history)
    except Exception as e:  # AI lỗi: vẫn trả điều luật, nhưng nói rõ lý do
        print(f"[legal_rag] LLM error: {e}")
        ai_error = str(e)
    if not text and not hits:
        return {"answer": "Mình đang hơi bận một chút, bạn thử gửi lại sau ít giây nhé!",
                "sources": [], "grounded": True, "ai_error": ai_error, "mode": r["mode"]}
    check = (
        verify_answer(text, hits)
        if text and hits
        else {"cited": [], "unverified_articles": [], "grounded": False}
    )
    return {
        "ai_error": ai_error,
        "answer": text or _fallback(hits),
        "grounded": check["grounded"] if (text and hits) else True,
        "unverified_articles": check["unverified_articles"],
        "mode": r["mode"],
        "confidence": r["confidence"],
        "query": r["query"] if r["query"] != question else None,
        "sources": [
            {
                "n": i, "law": h["law"], "article": h["article"], "clause": h["clause"],
                "title": h["title"], "chapter": h["chapter"], "excerpt": h["text"][:420],
                "cited": i in check["cited"], "score": h["score"], "similarity": h["dense"],
            }
            for i, h in enumerate(hits, 1)
            if not text or i in check["cited"]
        ],
    }


load()
